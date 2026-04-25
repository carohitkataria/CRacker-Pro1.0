"""CRacker Pro backend pytest suite."""
import os
import io
import time
import pytest
import requests

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://auto-fill-projects.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"
ADMIN_EMAIL = "admin@crackerpro.com"
ADMIN_PASSWORD = "Admin@123"


# ---------------- fixtures ----------------
@pytest.fixture(scope="session")
def admin_token():
    r = requests.post(f"{API}/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD}, timeout=15)
    if r.status_code == 429:
        # locked from prior run; reset by sleeping or hitting db is not possible from here
        pytest.skip("Admin locked due to brute-force; reset login_attempts in mongo")
    assert r.status_code == 200, r.text
    data = r.json()
    assert "access_token" in data and "user" in data
    assert data["user"]["email"] == ADMIN_EMAIL
    return data["access_token"]


@pytest.fixture(scope="session")
def admin_headers(admin_token):
    return {"Authorization": f"Bearer {admin_token}"}


# ---------------- AUTH ----------------
class TestAuth:
    def test_login_admin_success(self):
        r = requests.post(f"{API}/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD}, timeout=15)
        assert r.status_code == 200
        d = r.json()
        assert d["user"]["role"] == "admin"
        assert isinstance(d["access_token"], str) and len(d["access_token"]) > 20
        # bcrypt hash format check is implicit (login works) - explicit DB check skipped.

    def test_login_invalid_password(self):
        r = requests.post(f"{API}/auth/login", json={"email": ADMIN_EMAIL, "password": "WrongPass"}, timeout=15)
        assert r.status_code in (401, 429)

    def test_me_endpoint(self, admin_headers):
        r = requests.get(f"{API}/auth/me", headers=admin_headers, timeout=10)
        assert r.status_code == 200
        assert r.json()["email"] == ADMIN_EMAIL

    def test_me_unauthenticated(self):
        r = requests.get(f"{API}/auth/me", timeout=10)
        assert r.status_code == 401


# ---------------- DASHBOARD ----------------
class TestDashboard:
    def test_summary_structure(self, admin_headers):
        r = requests.get(f"{API}/dashboard/summary", headers=admin_headers, timeout=15)
        assert r.status_code == 200
        d = r.json()
        for k in ["totals", "stage_summary", "top_customers", "vendor_exposure",
                  "monthly_billing", "recognized_unbilled", "approvals_pending"]:
            assert k in d, f"missing key {k}"
        assert isinstance(d["stage_summary"], list) and len(d["stage_summary"]) == 5
        assert isinstance(d["monthly_billing"], list) and len(d["monthly_billing"]) == 6
        assert d["totals"]["total_projects"] >= 6


# ---------------- MASTERS ----------------
class TestMasters:
    def test_list_customers(self, admin_headers):
        r = requests.get(f"{API}/customers", headers=admin_headers, timeout=10)
        assert r.status_code == 200
        assert len(r.json()) >= 4

    def test_list_employees(self, admin_headers):
        r = requests.get(f"{API}/employees", headers=admin_headers, timeout=10)
        assert r.status_code == 200
        assert len(r.json()) >= 2

    def test_list_suppliers(self, admin_headers):
        r = requests.get(f"{API}/suppliers", headers=admin_headers, timeout=10)
        assert r.status_code == 200
        assert len(r.json()) >= 3

    def test_customer_crud(self, admin_headers):
        payload = {"customer_name": "TEST_Customer_QA", "sap_customer_code": "TESTC001",
                   "balance_outstanding_sap": 100, "country": "India"}
        r = requests.post(f"{API}/customers", headers=admin_headers, json=payload, timeout=10)
        assert r.status_code == 200, r.text
        cid = r.json()["id"]
        # update
        payload["risk_notes"] = "Updated"
        r = requests.put(f"{API}/customers/{cid}", headers=admin_headers, json=payload, timeout=10)
        assert r.status_code == 200 and r.json()["risk_notes"] == "Updated"
        # delete
        r = requests.delete(f"{API}/customers/{cid}", headers=admin_headers, timeout=10)
        assert r.status_code == 200


# ---------------- PROJECTS ----------------
class TestProjects:
    def test_list_projects_seeded(self, admin_headers):
        r = requests.get(f"{API}/projects", headers=admin_headers, timeout=10)
        assert r.status_code == 200
        rows = r.json()
        assert len(rows) >= 6
        # cache for other tests via attribute
        TestProjects.rows = rows

    def test_filter_by_stage_and_search(self, admin_headers):
        r = requests.get(f"{API}/projects", headers=admin_headers, params={"stage": "Pipeline"}, timeout=10)
        assert r.status_code == 200
        for p in r.json():
            assert p["current_stage"] == "Pipeline"
        r = requests.get(f"{API}/projects", headers=admin_headers, params={"search": "Smart"}, timeout=10)
        assert r.status_code == 200
        assert all("smart" in (p["project_name"] + p.get("wbs_element", "")).lower() for p in r.json())

    def test_get_one_project(self, admin_headers):
        rows = requests.get(f"{API}/projects", headers=admin_headers, timeout=10).json()
        pid = rows[0]["id"]
        r = requests.get(f"{API}/projects/{pid}", headers=admin_headers, timeout=10)
        assert r.status_code == 200
        assert r.json()["id"] == pid

    def test_create_update_project(self, admin_headers):
        custs = requests.get(f"{API}/customers", headers=admin_headers, timeout=10).json()
        payload = {
            "project_name": "TEST_QA_Project", "wbs_element": "WTEST.001",
            "customer_id": custs[0]["id"], "customer_name": custs[0]["customer_name"],
            "customer_po_number": "PO-TEST-1", "billing_type": "Milestone",
            "currency": "INR", "po_value": 1000000, "revenue_total": 1000000,
            "cost_total": 500000, "country": "India",
        }
        r = requests.post(f"{API}/projects", headers=admin_headers, json=payload, timeout=10)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["margin_pct"] == pytest.approx(50.0, rel=0.01)
        assert d["current_stage"] == "Pipeline"
        pid = d["id"]
        # update with cost increase → margin should drop
        payload["cost_total"] = 900000
        r = requests.put(f"{API}/projects/{pid}", headers=admin_headers, json=payload, timeout=10)
        assert r.status_code == 200
        assert r.json()["margin_pct"] == pytest.approx(10.0, rel=0.01)
        # cleanup
        requests.delete(f"{API}/projects/{pid}", headers=admin_headers, timeout=10)

    def test_transition_creates_low_margin_approval(self, admin_headers):
        rows = requests.get(f"{API}/projects", headers=admin_headers, timeout=10).json()
        # Find low-margin project that can transition to Customer PO
        # Seed has "Smart Buggy Management Software" at stage "Deal P&L" with ~10% margin
        target = None
        for p in rows:
            if (p.get("margin_pct", 100) < 15
                    and p.get("current_stage") == "Deal P&L"
                    and not p.get("project_name", "").startswith("TEST_")):
                target = p
                break
        if not target:
            pytest.skip("No low-margin Deal P&L project to test approval rule")
        pid = target["id"]
        r = requests.post(f"{API}/projects/{pid}/transition", headers=admin_headers,
                          json={"target_stage": "Customer PO"}, timeout=10)
        assert r.status_code == 200, r.text
        proj = r.json()
        # Approval should be Pending; stage should NOT have advanced yet
        assert proj["approval_status"] == "Pending", f"Expected Pending, got {proj['approval_status']}"
        assert proj["current_stage"] == "Deal P&L", \
            f"Stage should remain Deal P&L until approval, got {proj['current_stage']}"
        reqs = requests.get(f"{API}/approvals/requests", headers=admin_headers,
                            params={"status": "Pending"}, timeout=10).json()
        assert any(rq["project_id"] == pid and rq["target_stage"] == "Customer PO" for rq in reqs), \
            "Low-margin rule did not produce Pending approval for Customer PO"


# ---------------- REVENUE / COST LINES ----------------
class TestLines:
    def test_revenue_and_cost_lines(self, admin_headers):
        rows = requests.get(f"{API}/projects", headers=admin_headers, timeout=10).json()
        pid = rows[0]["id"]
        r = requests.get(f"{API}/projects/{pid}/revenue", headers=admin_headers, timeout=10)
        assert r.status_code == 200
        r = requests.get(f"{API}/projects/{pid}/cost", headers=admin_headers, timeout=10)
        assert r.status_code == 200
        # add new revenue line
        rl = {"revenue_code": "TEST-R", "description": "Test", "amount": 12345,
              "recognition_date": "2026-01-15", "is_billed": False}
        r = requests.post(f"{API}/projects/{pid}/revenue", headers=admin_headers, json=rl, timeout=10)
        assert r.status_code == 200, r.text
        rid = r.json()["id"]
        requests.delete(f"{API}/revenue/{rid}", headers=admin_headers, timeout=10)


# ---------------- EXCEL UPLOAD ----------------
class TestExcel:
    def test_template_download(self, admin_headers):
        r = requests.get(f"{API}/uploads/template/customer", headers=admin_headers, timeout=15)
        assert r.status_code == 200
        assert "spreadsheet" in r.headers.get("content-type", "")
        assert len(r.content) > 100

    def test_upload_customer_excel(self, admin_headers):
        # Get template, fill rows, upload
        tmpl = requests.get(f"{API}/uploads/template/customer", headers=admin_headers, timeout=15).content
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(tmpl))
            ws = wb.active
            # row 1 = headers; add 2 valid + 1 invalid (missing required)
            ws.append(["TEST_UPL_Cust1", "TUPL001", 1000, "", "Anil", "anil@test.com", "+91", "India"])
            ws.append(["TEST_UPL_Cust2", "TUPL002", 2000, "", "Sunil", "sunil@test.com", "+91", "India"])
            ws.append(["", "", "", "", "", "", "", ""])  # invalid - empty required
            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
        except Exception as e:
            pytest.skip(f"openpyxl not available: {e}")
        files = {"file": ("upload.xlsx", buf.getvalue(),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(f"{API}/uploads/customer", headers=admin_headers, files=files, timeout=20)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["success_rows"] >= 1
        # logs
        r = requests.get(f"{API}/uploads/logs", headers=admin_headers, timeout=10)
        assert r.status_code == 200 and len(r.json()) >= 1

    def test_upload_with_numeric_phone_and_listing(self, admin_headers):
        """Regression: openpyxl reads numeric phone cells as int. Backend must coerce
        to str so subsequent GET /api/customers does not fail with 500."""
        try:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.append(["customer_name", "sap_customer_code", "balance_outstanding_sap",
                       "risk_notes", "contact_person", "email", "phone", "country"])
            # phone is INT (not string) - the actual production hazard
            ws.append(["TEST_NumPhone_A", "TNP001", 500, "", "Ravi", "r@test.com", 9876543210, "India"])
            ws.append(["TEST_NumPhone_B", "TNP002", 600, "", "Asha", "a@test.com", 919999988888, "India"])
            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
        except Exception as e:
            pytest.skip(f"openpyxl not available: {e}")
        files = {"file": ("num_phone.xlsx", buf.getvalue(),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(f"{API}/uploads/customer", headers=admin_headers, files=files, timeout=20)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["success_rows"] >= 2, f"Expected >=2 rows ingested, got {d}"

        # Critical assertion: listing must still work (no 500 from Pydantic on int phone)
        r = requests.get(f"{API}/customers", headers=admin_headers, timeout=10)
        assert r.status_code == 200, f"GET /api/customers broke after numeric phone upload: {r.status_code} {r.text[:200]}"
        rows = r.json()
        # Find our uploads and verify phone is now stored as str
        ours = [c for c in rows if c.get("customer_name", "").startswith("TEST_NumPhone")]
        assert len(ours) >= 2, "Uploaded customers not found in listing"
        for c in ours:
            assert isinstance(c.get("phone"), str), f"phone should be str, got {type(c.get('phone'))}: {c.get('phone')}"
        # cleanup
        for c in ours:
            requests.delete(f"{API}/customers/{c['id']}", headers=admin_headers, timeout=10)


# ---------------- APPROVALS ----------------
class TestApprovals:
    def test_list_rules(self, admin_headers):
        r = requests.get(f"{API}/approvals/rules", headers=admin_headers, timeout=10)
        assert r.status_code == 200
        rules = r.json()
        assert len(rules) >= 2

    def test_rule_crud(self, admin_headers):
        payload = {"name": "TEST_Rule", "business_category": "Any",
                   "min_revenue": 1, "max_revenue": 10, "target_stage": "Deal P&L",
                   "approver_emails": [ADMIN_EMAIL], "approver_role": "admin", "is_active": True}
        r = requests.post(f"{API}/approvals/rules", headers=admin_headers, json=payload, timeout=10)
        assert r.status_code == 200, r.text
        rid = r.json()["id"]
        payload["name"] = "TEST_Rule_Upd"
        r = requests.put(f"{API}/approvals/rules/{rid}", headers=admin_headers, json=payload, timeout=10)
        assert r.status_code == 200 and r.json()["name"] == "TEST_Rule_Upd"
        r = requests.delete(f"{API}/approvals/rules/{rid}", headers=admin_headers, timeout=10)
        assert r.status_code == 200

    def test_approve_pending_request(self, admin_headers):
        # Find any pending request (created by prior tests or seeds)
        reqs = requests.get(f"{API}/approvals/requests", headers=admin_headers,
                            params={"status": "Pending"}, timeout=10).json()
        if not reqs:
            pytest.skip("No pending requests to approve")
        req_id = reqs[0]["id"]
        proj_id = reqs[0]["project_id"]
        target = reqs[0]["target_stage"]
        r = requests.post(f"{API}/approvals/requests/{req_id}/action", headers=admin_headers,
                          json={"action": "approve", "comment": "ok"}, timeout=10)
        assert r.status_code == 200, r.text
        # Confirm project moved to target stage
        proj = requests.get(f"{API}/projects/{proj_id}", headers=admin_headers, timeout=10).json()
        assert proj["current_stage"] == target


# ---------------- AUDIT ----------------
class TestAudit:
    def test_audit_listing(self, admin_headers):
        r = requests.get(f"{API}/audit", headers=admin_headers, timeout=10)
        assert r.status_code == 200
        assert isinstance(r.json(), list)

    def test_audit_filter_entity(self, admin_headers):
        r = requests.get(f"{API}/audit", headers=admin_headers, params={"entity_type": "project"}, timeout=10)
        assert r.status_code == 200
        for row in r.json():
            assert row["entity_type"] == "project"


# ---------------- ADMIN USER MGMT ----------------
class TestAdminUsers:
    def test_create_user_and_reset(self, admin_headers):
        ts = int(time.time())
        email = f"TEST_user_{ts}@cp.com"
        payload = {"email": email, "password": "Pass@12345", "name": "Test User",
                   "role": "finance", "location": "Delhi"}
        r = requests.post(f"{API}/admin/users", headers=admin_headers, json=payload, timeout=10)
        assert r.status_code == 200, r.text
        uid = r.json()["id"]
        assert r.json()["email"] == email.lower()
        # reset password
        r = requests.post(f"{API}/admin/users/reset-password", headers=admin_headers,
                          json={"user_id": uid, "new_password": "NewPass@123"}, timeout=10)
        assert r.status_code == 200
        # login with new password
        r = requests.post(f"{API}/auth/login",
                          json={"email": email.lower(), "password": "NewPass@123"}, timeout=10)
        assert r.status_code == 200
        finance_token = r.json()["access_token"]
        # finance user cannot list admin users
        r = requests.get(f"{API}/admin/users", headers={"Authorization": f"Bearer {finance_token}"}, timeout=10)
        assert r.status_code == 403
        # cleanup: deactivate
        requests.delete(f"{API}/admin/users/{uid}", headers=admin_headers, timeout=10)


# ---------------- BRUTE FORCE ----------------
class TestBruteForce:
    def test_lockout_after_repeated_attempts(self):
        # email-only lockout key now (k8s ingress rotates IPs). Should trigger by 6th attempt.
        email = f"locktest_{int(time.time())}@cp.com"
        statuses = []
        for _ in range(8):
            r = requests.post(f"{API}/auth/login",
                              json={"email": email, "password": "wrong"}, timeout=10)
            statuses.append(r.status_code)
            if r.status_code == 429:
                break
        assert 429 in statuses, f"429 never triggered in 8 attempts; statuses={statuses}"
        # Lockout should have triggered by attempt 6 (5 fails -> 6th sees locked_until)
        first_429 = statuses.index(429) + 1
        assert first_429 <= 6, f"Expected 429 by attempt 6 with email-only key, got at attempt {first_429}; statuses={statuses}"
